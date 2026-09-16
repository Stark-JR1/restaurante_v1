import io
import unittest
from datetime import datetime, date

from openpyxl import Workbook

from app.parsers.excel_parser import parse_excel_refeicoes
from app.services.auditoria_service import calcular_financeiro, validar_autorizacao_refeicao


class PlanilhaV2Test(unittest.TestCase):
    def criar_planilha(self):
        wb = Workbook()
        wb.remove(wb.active)

        cfg = wb.create_sheet('CONFIGURACAO')
        cfg.append(['CONFIGURAÇÃO DA MEDIÇÃO DE REFEIÇÕES'])
        cfg.append([])
        cfg.append(['PARÂMETRO', 'VALOR'])
        cfg.append(['Data inicial da medição', datetime(2026, 4, 1)])
        cfg.append(['Data final da medição', '31-4-26'])
        cfg.append([])
        cfg.append(['Café', 9])
        cfg.append(['Almoço', 29.90])
        cfg.append(['Janta', 30])

        col = wb.create_sheet('COLABORADORES')
        col.append(['CADASTRO'])
        col.append([])
        col.append(['MATRÍCULA', 'COLABORADOR', 'FUNÇÃO'])
        col.append([100, 'ANA TESTE', 'ADMIN'])
        col.append([200, 'BRUNO TESTE', 'OPERADOR'])

        exc = wb.create_sheet('EXCECOES')
        exc.append(['AUTORIZAÇÕES / EXCEÇÕES'])
        exc.append([])
        exc.append(['MATRÍCULA', 'COLABORADOR', 'INÍCIO', 'FIM', 'REFEIÇÃO', 'MOTIVO', 'AUTORIZADO', 'RESPONSÁVEL', 'STATUS'])
        exc.append([100, 'ANA TESTE', datetime(2026, 4, 1), datetime(2026, 4, 30), 'TODAS', 'TESTE', 'S', '', 'ATIVA'])

        lanc = wb.create_sheet('LANCAMENTOS')
        lanc.append(['LANÇAMENTOS'])
        lanc.append(['Nº CONTROLE', 'DATA', 'MATRÍCULA', 'COLABORADOR', 'TIPO', 'QUANTIDADE', 'VALOR UNIT.', 'TOTAL', 'EXCEÇÃO?', 'STATUS', 'OBSERVAÇÃO'])
        lanc.append(['CTRL-100-C', datetime(2026, 4, 2), 100, 'ANA TESTE', 'CAFÉ', 1, 9, 9, 'NÃO', 'OK', None])
        lanc.append(['CTRL-200-A', datetime(2026, 4, 3), 200, 'BRUNO TESTE', 'ALMOÇO', 2, 29.90, 59.80, 'NÃO', 'OK', None])

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def test_le_planilha_nova_e_corrige_fim_do_mes(self):
        dados = parse_excel_refeicoes(self.criar_planilha())
        self.assertEqual(dados['periodo_inicio'], date(2026, 4, 1))
        self.assertEqual(dados['periodo_fim'], date(2026, 4, 30))
        self.assertEqual(dados['formato_planilha'], 'v2_lancamentos')
        self.assertEqual(len(dados['refeicoes']), 2)
        self.assertEqual(sum(r['quantidade'] for r in dados['refeicoes']), 3)
        self.assertEqual(dados['colaboradores'][100], 'ANA TESTE')

    def test_precos_variaveis_sao_calculados_por_tipo(self):
        dados = parse_excel_refeicoes(self.criar_planilha())
        financeiro = calcular_financeiro(dados, [])
        self.assertTrue(financeiro['valor_unitario_variavel'])
        self.assertIsNone(financeiro['valor_unitario'])
        self.assertAlmostEqual(financeiro['total_provisionado_calculado'], 68.80, places=2)
        self.assertAlmostEqual(financeiro['valores_unitarios_por_tipo']['CAFÉ'], 9.0, places=2)

    def test_excecao_todas_autoriza_qualquer_refeicao(self):
        dados = parse_excel_refeicoes(self.criar_planilha())
        aut = validar_autorizacao_refeicao(dados['excecoes'], 100, date(2026, 4, 10), 'JANTA')
        self.assertEqual(aut['status'], 'AUTORIZADO')


if __name__ == '__main__':
    unittest.main()
