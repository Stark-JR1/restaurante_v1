import io
import re
import warnings
from datetime import date, datetime

import openpyxl

from app.utils.helpers import (
    first_money_after,
    first_text_after,
    quantidade_marcacao_refeicao,
    limpar_mojibake,
    normalizar_autorizado,
    normalizar_texto,
    normalizar_tipo_refeicao,
    parse_date_value,
    parse_money,
    tipo_refeicao_norm,
)


TIPOS_REFEICAO_DISPLAY = {
    "ALMOCO": "ALMO\u00c7O",
    "CAFE": "CAF\u00c9",
    "JANTA": "JANTA",
}


# ---------------------------------------------------------------------------
# Leitor da planilha nova (CONFIGURACAO + LANCAMENTOS)
# ---------------------------------------------------------------------------

def _resultado_base() -> dict:
    return {
        'periodo_inicio': None,
        'periodo_fim': None,
        'refeicoes': [],
        'linhas_refeicoes': [],
        'excecoes': [],
        'colaboradores': {},
        'restaurante': {
            'codigo': '',
            'nome': '',
            'razao_social': '',
            'cnpj': '',
            'cidade': ''
        },
        'totais_planilha': {
            'total_mensal_refeicoes': None,
            'valor_unitario': None,
            'valor_total': None,
            'valores_unitarios_por_tipo': {},
            'totais_por_tipo': {},
            'total_diario_por_dia': {},
            'total_diario_divergencias': []
        },
        'alertas_estrutura': [],
        'formato_planilha': 'v2_lancamentos',
        'aba_origem': 'LANCAMENTOS',
    }


def _valor_apos_rotulo(row, idx):
    for value in row[idx + 1:]:
        if value is not None and str(value).strip() != '':
            return value
    return None


def _parse_data_flexivel(valor, referencia: date | None = None):
    """Converte datas da planilha nova e corrige dia 31 em meses mais curtos.

    A planilha atual permite que a data final seja digitada como ``31-4-26``.
    Como abril termina no dia 30, o leitor interpreta esse caso como o último
    dia real do mês em vez de abortar a auditoria.
    """
    parsed = parse_date_value(valor)
    if parsed:
        return parsed

    if isinstance(valor, str):
        texto = valor.strip()
        # ISO gerado por alguns leitores: 2026-04-01T00:00:00.000Z
        m_iso = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})(?:[T\s].*)?$', texto)
        if m_iso:
            ano, mes, dia = map(int, m_iso.groups())
            try:
                return date(ano, mes, dia)
            except ValueError:
                return None

        m_br = re.match(r'^(\d{1,2})[\-/\.](\d{1,2})[\-/\.](\d{2}|\d{4})$', texto)
        if m_br:
            dia, mes, ano = map(int, m_br.groups())
            if ano < 100:
                ano += 2000
            try:
                return date(ano, mes, dia)
            except ValueError:
                # Se o único problema for um dia maior que o último dia do mês,
                # ajusta para o fechamento mensal real.
                try:
                    import calendar
                    ultimo = calendar.monthrange(ano, mes)[1]
                    if dia > ultimo and dia <= 31:
                        if referencia is None or (referencia.year == ano and referencia.month == mes):
                            return date(ano, mes, ultimo)
                except (ValueError, TypeError):
                    return None
    return None


def _ultimo_dia_mes(d: date) -> date:
    import calendar
    return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def _encontrar_cabecalho(ws, grupos, limite=20):
    """Retorna (índice 1-based, valores normalizados) do cabeçalho encontrado."""
    for numero_linha, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, limite), values_only=True), start=1):
        norm = [normalizar_texto(v) for v in row]
        if all(any(any(padrao in cel for padrao in grupo) for cel in norm) for grupo in grupos):
            return numero_linha, norm
    return None, []


def _indice_coluna(headers, *patterns, default=None):
    for i, header in enumerate(headers):
        if any(pattern in header for pattern in patterns):
            return i
    return default


def _tipos_v2(valor) -> list[str]:
    texto = normalizar_texto(valor)
    if texto in {'TODAS', 'TODOS', 'GERAL', 'TODAS AS REFEICOES', 'TODAS REFEICOES'}:
        return ['ALMOCO', 'CAFE', 'JANTA']
    tipos = normalizar_tipo_refeicao(valor)
    if tipos:
        return tipos
    tipo = tipo_refeicao_norm(valor)
    if tipo == 'ALMOÇO':
        return ['ALMOCO']
    if tipo == 'CAFÉ':
        return ['CAFE']
    if tipo == 'JANTA':
        return ['JANTA']
    return []


def _parse_excel_refeicoes_v2(wb) -> dict:
    result = _resultado_base()

    # 1) CONFIGURACAO: período e preços por tipo de refeição.
    ws_cfg = wb['CONFIGURACAO']
    precos = {}
    data_inicio_bruta = None
    data_fim_bruta = None

    for row in ws_cfg.iter_rows(values_only=True):
        norm = [normalizar_texto(v) for v in row]
        for idx, rotulo in enumerate(norm):
            if not rotulo:
                continue
            valor = _valor_apos_rotulo(row, idx)
            if rotulo.startswith('DATA') and ('INICIAL' in rotulo or 'INICIO' in rotulo):
                data_inicio_bruta = valor
            elif rotulo.startswith('DATA') and ('FINAL' in rotulo or 'FIM' in rotulo):
                data_fim_bruta = valor
            elif rotulo == 'CAFE':
                preco = parse_money(valor)
                if preco is not None:
                    precos['CAFE'] = preco
            elif 'ALMOCO' in rotulo:
                preco = parse_money(valor)
                if preco is not None:
                    precos['ALMOCO'] = preco
            elif 'JANTA' in rotulo:
                preco = parse_money(valor)
                if preco is not None:
                    precos['JANTA'] = preco

    periodo_inicio = _parse_data_flexivel(data_inicio_bruta)
    periodo_fim = _parse_data_flexivel(data_fim_bruta, periodo_inicio)

    # 2) COLABORADORES: detecta a linha de cabeçalho, sem depender de posição fixa.
    if 'COLABORADORES' in wb.sheetnames:
        ws_col = wb['COLABORADORES']
        header_col, headers_col = _encontrar_cabecalho(
            ws_col,
            [('MATRICULA',), ('COLABORADOR', 'NOME')],
            limite=15,
        )
        if header_col:
            mat_idx = _indice_coluna(headers_col, 'MATRICULA', default=0)
            nome_idx = _indice_coluna(headers_col, 'COLABORADOR', 'NOME', default=1)
            for row in ws_col.iter_rows(min_row=header_col + 1, values_only=True):
                mat = row[mat_idx] if mat_idx < len(row) else None
                if isinstance(mat, (int, float)):
                    nome = row[nome_idx] if nome_idx < len(row) else None
                    result['colaboradores'][int(mat)] = limpar_mojibake(str(nome).strip()) if nome else ''
        else:
            result['alertas_estrutura'].append('Cabeçalho da aba COLABORADORES não identificado.')
    else:
        result['alertas_estrutura'].append('Aba COLABORADORES não encontrada.')

    # 3) LANCAMENTOS: esta é a fonte oficial da nova versão.
    ws_lanc = wb['LANCAMENTOS']
    header_lanc, headers = _encontrar_cabecalho(
        ws_lanc,
        [('CONTROLE',), ('DATA',), ('MATRICULA',), ('TIPO',)],
        limite=15,
    )
    if not header_lanc:
        raise ValueError('Cabeçalho obrigatório da aba LANCAMENTOS não encontrado.')

    controle_idx = _indice_coluna(headers, 'CONTROLE')
    data_idx = _indice_coluna(headers, 'DATA')
    mat_idx = _indice_coluna(headers, 'MATRICULA')
    nome_idx = _indice_coluna(headers, 'COLABORADOR', 'NOME')
    tipo_idx = _indice_coluna(headers, 'TIPO')
    qtd_idx = _indice_coluna(headers, 'QUANTIDADE', 'QTD')
    unit_idx = _indice_coluna(headers, 'VALOR UNIT', 'UNIT')
    total_idx = _indice_coluna(headers, 'TOTAL')

    datas_lancadas = []
    total_valor = 0
    total_qtd = 0
    totais_por_tipo = {tipo: {'quantidade': 0, 'valor': 0} for tipo in ('CAFE', 'ALMOCO', 'JANTA')}
    total_diario = {}

    for numero_linha, row in enumerate(ws_lanc.iter_rows(min_row=header_lanc + 1, values_only=True), start=header_lanc + 1):
        if not row or not any(v is not None and str(v).strip() != '' for v in row):
            continue

        mat_val = row[mat_idx] if mat_idx is not None and mat_idx < len(row) else None
        data_val = row[data_idx] if data_idx is not None and data_idx < len(row) else None
        tipo_val = row[tipo_idx] if tipo_idx is not None and tipo_idx < len(row) else None
        qtd_val = row[qtd_idx] if qtd_idx is not None and qtd_idx < len(row) else 1

        # Linhas vazias com fórmulas sem resultado não viram lançamento.
        if not isinstance(mat_val, (int, float)) or not tipo_val or data_val in (None, ''):
            continue

        data_ref = _parse_data_flexivel(data_val, periodo_inicio)
        if not data_ref:
            result['alertas_estrutura'].append(f'LANCAMENTOS linha {numero_linha}: data inválida ou não reconhecida.')
            continue

        tipos = _tipos_v2(tipo_val)
        if not tipos:
            result['alertas_estrutura'].append(f'LANCAMENTOS linha {numero_linha}: tipo de refeição não reconhecido: {tipo_val}.')
            continue

        quantidade = quantidade_marcacao_refeicao(qtd_val)
        if quantidade <= 0:
            continue

        matricula = int(mat_val)
        nome_val = row[nome_idx] if nome_idx is not None and nome_idx < len(row) else None
        nome = limpar_mojibake(str(nome_val).strip()) if nome_val else result['colaboradores'].get(matricula, '')
        controle = row[controle_idx] if controle_idx is not None and controle_idx < len(row) else None
        controle = limpar_mojibake(str(controle).strip()) if controle else ''
        unit_linha = parse_money(row[unit_idx]) if unit_idx is not None and unit_idx < len(row) else None
        total_linha = parse_money(row[total_idx]) if total_idx is not None and total_idx < len(row) else None

        datas_lancadas.append(data_ref)
        total_calculado_linha = 0
        valor_calculado_linha = 0

        for tipo in tipos:
            valor_unitario = unit_linha if len(tipos) == 1 and unit_linha is not None else precos.get(tipo)
            if valor_unitario is None and unit_linha is not None:
                valor_unitario = unit_linha

            display = TIPOS_REFEICAO_DISPLAY.get(tipo, tipo)
            valor_total_ref = (valor_unitario * quantidade) if valor_unitario is not None else None
            result['refeicoes'].append({
                'matricula': matricula,
                'nome': nome,
                'tipo': display,
                'tipo_original': limpar_mojibake(str(tipo_val).strip()),
                'dia': data_ref,
                'numero_dia': data_ref.day,
                'quantidade': quantidade,
                'linha_excel': numero_linha,
                'numero_controle': controle,
                'valor_unitario': valor_unitario,
                'valor_total': valor_total_ref,
            })

            total_calculado_linha += quantidade
            if valor_total_ref is not None:
                valor_calculado_linha += valor_total_ref
                total_valor += valor_total_ref
                totais_por_tipo.setdefault(tipo, {'quantidade': 0, 'valor': 0})['valor'] += valor_total_ref
            totais_por_tipo.setdefault(tipo, {'quantidade': 0, 'valor': 0})['quantidade'] += quantidade
            total_qtd += quantidade
            total_diario[data_ref.day] = total_diario.get(data_ref.day, 0) + quantidade

        # Na nova planilha QUANTIDADE é a referência da linha. Se houver tipo
        # composto, a auditoria expande um registro para cada tipo reconhecido.
        valor_total_informado = total_linha
        if valor_total_informado is None and valor_calculado_linha:
            valor_total_informado = valor_calculado_linha

        result['linhas_refeicoes'].append({
            'linha_excel': numero_linha,
            'matricula': matricula,
            'nome': nome,
            'tipo': ' / '.join(TIPOS_REFEICAO_DISPLAY.get(t, t) for t in tipos),
            'tipo_original': limpar_mojibake(str(tipo_val).strip()),
            'dias': [data_ref],
            'total_calculado': total_calculado_linha,
            # Evita comparar QUANTIDADE=1 contra dois tipos quando uma linha
            # composta é legitimamente expandida pelo leitor.
            'total_informado': total_calculado_linha,
            'quantidade_informada': quantidade,
            'valor_unitario': unit_linha if unit_linha is not None else (precos.get(tipos[0]) if len(tipos) == 1 else None),
            'valor_total_informado': valor_total_informado,
            'numero_controle': controle,
        })

    # Período: CONFIGURACAO manda. Se a data final for inválida/ausente, usa o
    # último dia real do mês inicial. Se a data inicial também faltar, cai para
    # as datas efetivamente lançadas.
    if periodo_inicio is None and datas_lancadas:
        periodo_inicio = min(datas_lancadas)
    if periodo_inicio is None:
        raise ValueError('Período não encontrado na aba CONFIGURACAO nem nos LANCAMENTOS.')
    if periodo_fim is None:
        periodo_fim = _ultimo_dia_mes(periodo_inicio)
    if periodo_fim < periodo_inicio:
        periodo_fim = _ultimo_dia_mes(periodo_inicio)

    result['periodo_inicio'] = periodo_inicio
    result['periodo_fim'] = periodo_fim

    # 4) Totais e preços. A versão nova possui preço independente para cada tipo.
    result['totais_planilha']['total_mensal_refeicoes'] = total_qtd
    result['totais_planilha']['valor_total'] = total_valor
    result['totais_planilha']['total_diario_por_dia'] = total_diario
    result['totais_planilha']['totais_por_tipo'] = {
        TIPOS_REFEICAO_DISPLAY.get(k, k): {
            'quantidade': int(v['quantidade']),
            'valor': v['valor'],
        }
        for k, v in totais_por_tipo.items()
    }
    result['totais_planilha']['valores_unitarios_por_tipo'] = {
        TIPOS_REFEICAO_DISPLAY.get(k, k): v for k, v in precos.items()
    }
    precos_unicos = {str(v) for v in precos.values()}
    if len(precos_unicos) == 1 and precos:
        result['totais_planilha']['valor_unitario'] = next(iter(precos.values()))

    # 5) EXCECOES: cabeçalho fica na linha 3 na nova planilha.
    aba_excecoes = next((nome for nome in wb.sheetnames if normalizar_texto(nome) in {'EXCECOES', 'AUTORIZACOES'}), None)
    if aba_excecoes:
        ws_exc = wb[aba_excecoes]
        header_exc, headers_exc = _encontrar_cabecalho(
            ws_exc,
            [('MATRICULA',), ('INICIO',), ('FIM',), ('REFEICAO', 'TIPO')],
            limite=15,
        )
        if header_exc:
            mat_e = _indice_coluna(headers_exc, 'MATRICULA', default=0)
            nome_e = _indice_coluna(headers_exc, 'COLABORADOR', 'NOME', default=1)
            ini_e = _indice_coluna(headers_exc, 'INICIO', default=2)
            fim_e = _indice_coluna(headers_exc, 'FIM', default=3)
            tipo_e = _indice_coluna(headers_exc, 'REFEICAO', 'TIPO', default=4)
            motivo_e = _indice_coluna(headers_exc, 'MOTIVO', 'JUSTIFICATIVA', default=5)
            aut_e = _indice_coluna(headers_exc, 'AUTORIZADO', default=6)
            resp_e = _indice_coluna(headers_exc, 'RESPONSAVEL')
            status_e = _indice_coluna(headers_exc, 'STATUS')

            for numero_linha, row in enumerate(ws_exc.iter_rows(min_row=header_exc + 1, values_only=True), start=header_exc + 1):
                mat_val = row[mat_e] if mat_e < len(row) else None
                if not isinstance(mat_val, (int, float)):
                    continue

                data_ini = _parse_data_flexivel(row[ini_e] if ini_e < len(row) else None, periodo_inicio)
                data_fim = _parse_data_flexivel(row[fim_e] if fim_e < len(row) else None, data_ini or periodo_inicio)
                tipo_bruto = row[tipo_e] if tipo_e < len(row) else None
                tipos = _tipos_v2(tipo_bruto)
                autorizado = normalizar_autorizado(row[aut_e] if aut_e < len(row) else None)
                status_registro = normalizar_texto(row[status_e]) if status_e is not None and status_e < len(row) else 'ATIVA'
                ativa = status_registro not in {'INATIVA', 'INATIVO', 'CANCELADA', 'CANCELADO'}
                valido = bool(data_ini and data_fim and tipos and ativa)

                result['excecoes'].append({
                    'matricula': int(mat_val),
                    'nome': limpar_mojibake(str(row[nome_e]).strip()) if nome_e < len(row) and row[nome_e] else '',
                    'data_inicio': data_ini,
                    'data_fim': data_fim,
                    'tipo': tipo_refeicao_norm(tipo_bruto or ''),
                    'tipos': tipos,
                    'motivo': limpar_mojibake(str(row[motivo_e]).strip()) if motivo_e < len(row) and row[motivo_e] else '',
                    'autorizado': autorizado,
                    'valido': valido,
                    'data_valida': data_ini is not None and data_fim is not None,
                    'linha_excel': numero_linha,
                    'responsavel': limpar_mojibake(str(row[resp_e]).strip()) if resp_e is not None and resp_e < len(row) and row[resp_e] else '',
                    'status_registro': status_registro,
                })
        else:
            result['alertas_estrutura'].append('Cabeçalho da aba EXCECOES não identificado.')
    else:
        result['alertas_estrutura'].append('Aba EXCECOES não encontrada.')

    return result


def parse_excel_refeicoes(xlsx_bytes: bytes) -> dict:
    """Lê a planilha de refeições.

    Prioriza o layout novo, baseado nas abas CONFIGURACAO e LANCAMENTOS.
    Mantém o leitor legado da aba REFEIÇÕES para não quebrar arquivos antigos.
    """
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='Data Validation extension is not supported.*')
        warnings.filterwarnings('ignore', message='Conditional Formatting extension is not supported.*')
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    nomes_norm = {normalizar_texto(nome): nome for nome in wb.sheetnames}
    if 'CONFIGURACAO' in nomes_norm and 'LANCAMENTOS' in nomes_norm:
        # O template oficial usa exatamente estes nomes; o mapeamento acima
        # também aceita variação de acentos/capitalização no reconhecimento.
        if nomes_norm['CONFIGURACAO'] != 'CONFIGURACAO' or nomes_norm['LANCAMENTOS'] != 'LANCAMENTOS':
            wb[nomes_norm['CONFIGURACAO']].title = 'CONFIGURACAO'
            wb[nomes_norm['LANCAMENTOS']].title = 'LANCAMENTOS'
        return _parse_excel_refeicoes_v2(wb)

    return _parse_excel_refeicoes_legacy(xlsx_bytes)


def _parse_excel_refeicoes_legacy(xlsx_bytes: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)

    result = {
        'periodo_inicio': None,
        'periodo_fim': None,
        'refeicoes': [],
        'linhas_refeicoes': [],
        'excecoes': [],
        'colaboradores': {},
        'restaurante': {
            'codigo': '',
            'nome': '',
            'razao_social': '',
            'cnpj': '',
            'cidade': ''
        },
        'totais_planilha': {
            'total_mensal_refeicoes': None,
            'valor_unitario': None,
            'valor_total': None,
            'total_diario_por_dia': {},
            'total_diario_divergencias': []
        },
        'alertas_estrutura': []
    }

    if 'COLABORADORES' in wb.sheetnames:
        ws = wb['COLABORADORES']
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and isinstance(row[0], (int, float)):
                result['colaboradores'][int(row[0])] = limpar_mojibake(str(row[1]).strip()) if row[1] else ''
    else:
        result['alertas_estrutura'].append("Aba COLABORADORES não encontrada.")

    if 'REFEIÇÕES' not in wb.sheetnames:
        raise ValueError("Aba obrigatória REFEIÇÕES não encontrada.")

    ws = wb['REFEIÇÕES']
    rows = list(ws.iter_rows(values_only=True))

    for row in rows:
        normalized_row = [normalizar_texto(c) for c in row]
        for idx, texto in enumerate(normalized_row):
            if texto.startswith("RESTAURANTE"):
                valor = first_text_after(row, idx)
                if valor:
                    valor_texto = limpar_mojibake(str(valor).strip())
                    result['restaurante']['nome'] = valor_texto
                    m_rest = re.match(r'^\s*(\d+)\s*-\s*(.+?)(?:\s+-\s+(.+))?\s*$', valor_texto)
                    if m_rest:
                        result['restaurante']['codigo'] = m_rest.group(1).strip()
                        result['restaurante']['razao_social'] = limpar_mojibake(m_rest.group(2).strip())
                        result['restaurante']['nome'] = limpar_mojibake((m_rest.group(3) or m_rest.group(2)).strip())
            elif texto.startswith("CNPJ"):
                valor = first_text_after(row, idx)
                if valor:
                    result['restaurante']['cnpj'] = limpar_mojibake(str(valor).strip())
            elif texto.startswith("CIDADE"):
                valor = first_text_after(row, idx)
                if valor:
                    result['restaurante']['cidade'] = limpar_mojibake(str(valor).strip())
            elif "TOTAL DE MENSAL" in texto or "TOTAL MENSAL" in texto:
                valor = first_money_after(row, idx)
                if valor is not None:
                    result['totais_planilha']['total_mensal_refeicoes'] = int(parse_money(valor))
            elif "VALOR UNIT" in texto:
                valor = first_money_after(row, idx)
                if valor is not None:
                    result['totais_planilha']['valor_unitario'] = parse_money(valor)
            elif texto == "VALOR TOTAL" or "VALOR TOTAL" in texto:
                valor = first_money_after(row, idx)
                if valor is not None:
                    result['totais_planilha']['valor_total'] = parse_money(valor)

    header_row = rows[0]
    for row in rows[:5]:
        for cell in row:
            texto = normalizar_texto(cell)
            if "DATA MEDICAO" in texto:
                continue
            data_cell = parse_date_value(cell)
            if data_cell:
                if result['periodo_inicio'] is None:
                    result['periodo_inicio'] = data_cell
                elif result['periodo_fim'] is None and data_cell != result['periodo_inicio']:
                    result['periodo_fim'] = data_cell

    for cell in header_row:
        if isinstance(cell, datetime):
            if result['periodo_inicio'] is None:
                result['periodo_inicio'] = cell.date()
            else:
                result['periodo_fim'] = cell.date()

    if not result['periodo_inicio']:
        raise ValueError("Período não encontrado na planilha REFEIÇÕES.")
    if not result['periodo_fim']:
        result['periodo_fim'] = result['periodo_inicio']

    ano = result['periodo_inicio'].year
    mes = result['periodo_inicio'].month

    header_idx = None
    for idx, row in enumerate(rows[:10]):
        normalized = [normalizar_texto(c) for c in row]
        if any("MATRICULA" in c for c in normalized) and any("TIPO" in c for c in normalized):
            header_idx = idx
            break

    if header_idx is None:
        header_idx = 2
        mat_idx, nome_idx, tipo_idx = 1, 2, 3
        dia_indices = {i + 1: 4 + i for i in range(31)}
        total_idx = 35
        valor_unit_idx = None
        valor_total_idx = None
    else:
        header_row_norm = [normalizar_texto(c) for c in rows[header_idx]]
        mat_idx = next((i for i, h in enumerate(header_row_norm) if "MATRICULA" in h), 1)
        nome_idx = next((i for i, h in enumerate(header_row_norm) if h == "NOME" or "COLABORADOR" in h), 2)
        tipo_idx = next((i for i, h in enumerate(header_row_norm) if "TIPO" in h), 3)
        dia_indices = {}
        for i, h in enumerate(header_row_norm):
            m = re.search(r"DIA\s*(\d{1,2})", h)
            if m:
                dia_indices[int(m.group(1))] = i
        if not dia_indices:
            dia_indices = {i + 1: 4 + i for i in range(31)}
        total_idx = next((i for i, h in enumerate(header_row_norm) if h == "TOTAL" or "TOTAL REFE" in h), None)
        valor_unit_idx = next((i for i, h in enumerate(header_row_norm) if "UNIT" in h or "VALOR UNIT" in h), None)
        valor_total_idx = next((i for i, h in enumerate(header_row_norm) if "VALOR TOTAL" in h or h == "TOTAL R$"), None)

    for numero_linha, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not row:
            continue

        mat_val = row[mat_idx] if mat_idx < len(row) else None
        nome_val = row[nome_idx] if nome_idx < len(row) else None
        tipo_val = row[tipo_idx] if tipo_idx < len(row) else None

        if not mat_val or not isinstance(mat_val, (int, float)):
            continue
        if not tipo_val:
            continue

        matricula = int(mat_val)
        nome = limpar_mojibake(str(nome_val).strip()) if nome_val else ""
        tipos_norm = normalizar_tipo_refeicao(tipo_val)
        if not tipos_norm:
            tipo_norm = tipo_refeicao_norm(tipo_val)
            tipos_norm = [tipo_norm] if tipo_norm else []
        total_calculado = 0
        dias_marcados = []

        for numero_dia, col_idx in sorted(dia_indices.items()):
            val = row[col_idx] if col_idx < len(row) else None
            quantidade = quantidade_marcacao_refeicao(val)
            if quantidade:
                try:
                    d = date(ano, mes, numero_dia)
                    total_calculado += quantidade * len(tipos_norm)
                    dias_marcados.append(d)
                    for tipo_refeicao in tipos_norm:
                        result['refeicoes'].append({
                            'matricula': matricula,
                            'nome': nome,
                            'tipo': TIPOS_REFEICAO_DISPLAY.get(tipo_refeicao, tipo_refeicao),
                            'tipo_original': limpar_mojibake(str(tipo_val).strip()),
                            'dia': d,
                            'numero_dia': numero_dia,
                            'quantidade': quantidade,
                            'linha_excel': numero_linha
                        })
                except ValueError:
                    result['alertas_estrutura'].append(
                        f"Linha {numero_linha}: dia {numero_dia} inválido para o mês da medição."
                    )

        total_informado = row[total_idx] if total_idx is not None and total_idx < len(row) else None
        valor_unitario = parse_money(row[valor_unit_idx]) if valor_unit_idx is not None and valor_unit_idx < len(row) else None
        valor_total = parse_money(row[valor_total_idx]) if valor_total_idx is not None and valor_total_idx < len(row) else None

        result['linhas_refeicoes'].append({
            'linha_excel': numero_linha,
            'matricula': matricula,
            'nome': nome,
            'tipo': " / ".join(TIPOS_REFEICAO_DISPLAY.get(tipo, tipo) for tipo in tipos_norm),
            'tipo_original': limpar_mojibake(str(tipo_val).strip()),
            'dias': dias_marcados,
            'total_calculado': total_calculado,
            'total_informado': int(total_informado) if isinstance(total_informado, (int, float)) else None,
            'valor_unitario': valor_unitario,
            'valor_total_informado': valor_total
        })

    total_por_dia_calculado = {}
    for ref in result['refeicoes']:
        total_por_dia_calculado[ref['numero_dia']] = (
            total_por_dia_calculado.get(ref['numero_dia'], 0) + ref.get('quantidade', 1)
        )

    for numero_linha, row in enumerate(rows, start=1):
        if not row:
            continue
        normalized = [normalizar_texto(c) for c in row]
        if not any("TOTAL DIARIO" in c for c in normalized):
            continue

        for numero_dia, col_idx in sorted(dia_indices.items()):
            informado = row[col_idx] if col_idx < len(row) else None
            if isinstance(informado, (int, float)):
                informado_int = int(informado)
                result['totais_planilha']['total_diario_por_dia'][numero_dia] = informado_int
                calculado = total_por_dia_calculado.get(numero_dia, 0)
                if informado_int != calculado:
                    result['totais_planilha']['total_diario_divergencias'].append({
                        'linha_excel': numero_linha,
                        'dia': numero_dia,
                        'informado': informado_int,
                        'calculado': calculado
                    })
        break

    aba_excecoes = next((nome for nome in wb.sheetnames if normalizar_texto(nome) in {"EXCECOES", "AUTORIZACOES"}), None)
    if aba_excecoes:
        ws_exc = wb[aba_excecoes]
        rows_exc = list(ws_exc.iter_rows(values_only=True))

        header_idx_exc = 0
        header_exc = [normalizar_texto(c) for c in rows_exc[0]] if rows_exc else []

        def find_col(*patterns, default=None):
            for i, header in enumerate(header_exc):
                if any(pattern in header for pattern in patterns):
                    return i
            return default

        mat_exc_idx = find_col("MATRICULA", default=0)
        nome_exc_idx = find_col("NOME", "COLABORADOR", default=1)
        data_ini_idx = find_col("DATA_INICIO", "DATA INICIO", default=2)
        data_fim_idx = find_col("DATA_FIM", "DATA FIM", default=3)
        tipo_exc_idx = find_col("TIPO", default=4)
        motivo_exc_idx = find_col("MOTIVO", "JUSTIFICATIVA", default=5)
        autorizado_idx = find_col("AUTORIZADO", default=6)

        for numero_linha, row in enumerate(rows_exc[header_idx_exc + 1:], start=header_idx_exc + 2):
            mat_val = row[mat_exc_idx] if row and mat_exc_idx < len(row) else None
            if not row or not mat_val:
                continue
            if not isinstance(mat_val, (int, float)):
                continue

            data_ini = parse_date_value(row[data_ini_idx] if data_ini_idx < len(row) else None)
            data_fim = parse_date_value(row[data_fim_idx] if data_fim_idx < len(row) else None)
            autorizado = normalizar_autorizado(row[autorizado_idx] if autorizado_idx < len(row) else None)
            tipos = normalizar_tipo_refeicao(row[tipo_exc_idx] if tipo_exc_idx < len(row) else None)
            valido = data_ini is not None and data_fim is not None and bool(tipos)

            nome_exc = limpar_mojibake(str(row[nome_exc_idx]).strip()) if nome_exc_idx < len(row) and row[nome_exc_idx] else ''
            motivo_exc = limpar_mojibake(str(row[motivo_exc_idx]).strip()) if motivo_exc_idx < len(row) and row[motivo_exc_idx] else ''

            result['excecoes'].append({
                'matricula': int(mat_val),
                'nome': nome_exc,
                'data_inicio': data_ini,
                'data_fim': data_fim,
                'tipo': tipo_refeicao_norm(row[tipo_exc_idx] if tipo_exc_idx < len(row) else ''),
                'tipos': tipos,
                'motivo': motivo_exc,
                'autorizado': autorizado,
                'valido': valido,
                'data_valida': data_ini is not None and data_fim is not None,
                'linha_excel': numero_linha
            })
    else:
        result['alertas_estrutura'].append("Aba EXCEÇÕES não encontrada.")

    return result
